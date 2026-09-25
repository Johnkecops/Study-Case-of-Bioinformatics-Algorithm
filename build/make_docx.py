# -*- coding: utf-8 -*-
"""Build IKSAD-Bioalgo-Manuscript-Arli.docx"""
import sys, re, json, copy
sys.path.insert(0, "build")
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement, parse_xml

import os
SKIP=set(os.environ.get("SKIP","").split(","))
import content as C, refs as R, tables as T, omml
from areas import AREAS

EQ = omml.build()
FONT, SIZE = "Times New Roman", Pt(12)
TEXTW = Cm(16.0)

# ---------- low-level helpers ----------
def style_run(run, size=SIZE, bold=False, italic=False, font=FONT):
    run.font.name = font; run.font.size = size; run.bold = bold; run.italic = italic
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a), font)
    return run

def para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=5, spacing=1.0, keep=False):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.alignment = align; pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = spacing; pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if keep: pf.keep_with_next = True
    return p

# in-text citation substitution: {PMID} parenthetical, [[PMID]] narrative
CIT = re.compile(r"\{(\d+)\}|\[\[(\d+)\]\]")
def add_text(p, text, size=SIZE, bold=False, italic=False):
    out, last = [], 0
    for m in CIT.finditer(text):
        out.append(text[last:m.start()])
        pmid = m.group(1) or m.group(2)
        out.append(R.cite(pmid, narrative=bool(m.group(2))))
        last = m.end()
    out.append(text[last:])
    s = "".join(out)
    s = re.sub(r"\)\s+\(", "; ", s)          # merge adjacent parenthetical citations
    style_run(p.add_run(s), size, bold, italic)
    return s

def heading(doc, text, level=1):
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=12 if level == 1 else 10, after=6, keep=True)
    style_run(p.add_run(text.upper() if level == 1 else text), SIZE, bold=True)
    return p

def shade(cell_or_para, hexcolor):
    el = cell_or_para._element
    pr = el.find(qn('w:tcPr')) if el.tag.endswith('}tc') else el.get_or_add_pPr()
    if pr is None: pr = OxmlElement('w:tcPr'); el.insert(0, pr)
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexcolor)
    pr.append(sh)

def fixed_layout(tbl, total_cm=16.0):
    pr = tbl._element.find(qn('w:tblPr'))
    if pr is None: pr = OxmlElement('w:tblPr'); tbl._element.insert(0, pr)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); pr.append(lay)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(int(total_cm * 567))); w.set(qn('w:type'), 'dxa'); pr.append(w)

def set_borders(tbl_or_cell, sz=6, color="000000"):
    el = tbl_or_cell._element
    tag = 'w:tcPr' if el.tag.endswith('}tc') else 'w:tblPr'
    pr = el.find(qn(tag))
    if pr is None: pr = OxmlElement(tag); el.insert(0, pr)
    borders = OxmlElement('w:tcBorders' if tag == 'w:tcPr' else 'w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if tag == 'w:tcPr' and edge.startswith('inside'): continue
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)
    pr.append(borders)

# ---------- document ----------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec, m, Cm(2.5))
n = doc.styles['Normal']; n.font.name = FONT; n.font.size = SIZE
n.paragraph_format.line_spacing = 1.0
n.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
n.paragraph_format.space_after = Pt(0)
n.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

# --- title ---
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=12)
style_run(p.add_run(C.TITLE), SIZE, bold=True)

# --- author block ---
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=2)
style_run(p.add_run(C.AUTHOR["name"]), SIZE, bold=True)
for line in C.AUTHOR["lines"]:
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=1)
    style_run(p.add_run(line), Pt(11))
para(doc, after=6)

# --- abstract ---
heading(doc, "Abstract")
p = para(doc); add_text(p, C.ABSTRACT)
p = para(doc, after=6)
style_run(p.add_run("Keywords: "), SIZE, bold=True); style_run(p.add_run(C.KEYWORDS), SIZE)

# --- introduction ---
heading(doc, "Introduction")
for t in C.INTRO:
    add_text(para(doc), t)

# --- methods ---
heading(doc, "Materials and Methods")
for sub, paras in C.METHODS:
    heading(doc, sub, level=2)
    for t in paras: add_text(para(doc), t)

# Table 1 + Figure 1 belong to Methods
def add_table(spec, widths):
    if "tbl" in SKIP: return
    cap = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, keep=True)
    style_run(cap.add_run(spec["caption"]), Pt(11), bold=True)
    tbl = doc.add_table(rows=1, cols=len(spec["head"]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_borders(tbl); fixed_layout(tbl, sum(c.cm for c in widths))
    for i, h in enumerate(spec["head"]):
        c = tbl.rows[0].cells[i]; c.text = ""
        pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(2); pp.paragraph_format.space_before = Pt(2)
        style_run(pp.add_run(h), Pt(10), bold=True); shade(c, "E8E8E8")
    for row in spec["rows"]:
        cells = tbl.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            pp.paragraph_format.space_after = Pt(2); pp.paragraph_format.space_before = Pt(2)
            bold = (row[0] == "Total")
            style_run(pp.add_run(v), Pt(10), bold=bold)
    for r in tbl.rows:
        for i, c in enumerate(r.cells): c.width = widths[i]
    para(doc, after=6)

def add_figure(path, caption, width_cm=15.2):
    if "fig" in SKIP: return
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
    p.add_run().add_picture(path, width=Cm(width_cm))
    cap = para(doc, WD_ALIGN_PARAGRAPH.LEFT, after=8)
    style_run(cap.add_run(caption), Pt(11))

add_table(T.TABLE1, [Cm(3.0), Cm(6.5), Cm(6.5)])
add_figure("figures/fig1_prisma.png",
           "Figure 1. PRISMA 2020 flow diagram for the review. Counts on the left arm are records retrieved from "
           "PubMed/MEDLINE on 4 September 2026; counts on the right arm are records identified through citation "
           "chaining and hand searching. Screening was capped at the 60 highest-ranking records per search block, "
           "so 18,690 unique records were not sought for retrieval.", 13.2)

# --- findings ---
heading(doc, "Findings and Discussion")
def eq_para(key, number):
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run("\t"))
    p._p.append(parse_xml(EQ[key]))
    style_run(p.add_run("\t(%d)" % number))

def box_para(num, title, code, note):
    if "box" in SKIP: return
    cap = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3, keep=True)
    style_run(cap.add_run("Box %d. %s" % (num, title)), Pt(11), bold=True)
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells[0]; c.width = Cm(16.0)
    set_borders(tbl, sz=8); fixed_layout(tbl, 16.0); shade(c, "F2F2F2")
    c.text = ""
    first = True
    for line in code.split("\n"):
        pp = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.space_before = Pt(0)
        pp.paragraph_format.line_spacing = 1.0
        pp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style_run(pp.add_run(line if line else " "), Pt(9.5), font="Courier New")
    nt = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=2, after=8)
    style_run(nt.add_run(note), Pt(10), italic=True)

FIGCAP = {2: ("figures/fig2_synthesis.png",
   "Figure 2. Shared structure of the three application areas. Each column follows one area from its training "
   "signal, through the objective function that ranks candidates, to the generated output and the experimental "
   "route by which that output can be verified. The middle row identifies the equation in the text that states "
   "each objective.", 15.2)}
TBLSPEC = {2: (T.TABLE2, [Cm(4.6), Cm(1.0), Cm(2.2), Cm(2.2), Cm(2.2), Cm(3.8)]),
           3: (T.TABLE3, [Cm(3.0), Cm(4.4), Cm(4.4), Cm(4.2)])}

for item in C.FINDINGS:
    k = item[0]
    if   k == "sub": heading(doc, item[1], level=2)
    elif k == "p":   add_text(para(doc), item[1])
    elif k == "eq":  eq_para(item[1], item[2])
    elif k == "box": box_para(item[1], item[2], item[3], item[4])
    elif k == "tbl": add_table(*TBLSPEC[item[1]])
    elif k == "fig": add_figure(*FIGCAP[item[1]])

# --- conclusion ---
heading(doc, "Conclusion and Recommendations")
for t in C.CONCLUSION: add_text(para(doc), t)

# --- references ---
heading(doc, "References")
for pmid in ([] if "ref" in SKIP else R.ORDER):
    e = R.entry(pmid)
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, after=3)
    pf = p.paragraph_format; pf.left_indent = Cm(1.27); pf.first_line_indent = Cm(-1.27)
    style_run(p.add_run("%s (%s). %s. " % (e["auth"], e["year"], e["title"])), Pt(11))
    style_run(p.add_run(e["journal"]), Pt(11), italic=True)
    if e["tail"]:
        vol = e["tail"].split("(")[0].split(",")[0]
        rest = e["tail"][len(vol):]
        style_run(p.add_run(", "), Pt(11))
        style_run(p.add_run(vol), Pt(11), italic=True)
        style_run(p.add_run(rest), Pt(11))
    style_run(p.add_run(". doi:%s" % e["doi"] if e["doi"] else "."), Pt(11))

OUT = os.environ.get("OUT","IKSAD-Bioalgo-Manuscript-Arli.docx")
doc.save(OUT)
print("saved", OUT)
